"""Deck parsing + slide-aware chunking — pure functions, no Prefect/DB/Qdrant.

Same contract as paper.py, one level up (fetch -> parse -> caption -> chunk ->
embed -> index; the Prefect flow lives in document_pipeline.py). A deck is
either a PDF (one page = one slide) or a PPTX (native slides, shapes, speaker
notes) — this module hides that behind one shape:

  parse_deck()    file -> [{slide, text, notes}], 1-indexed, one entry per
                  slide (including empty ones — numbering never drifts)
  needs_caption() a slide with little/no extracted text is probably a chart or
                  diagram drawn as vector shapes; caption it instead
  render_slides() slide numbers -> {slide: jpeg} — rasterized page (PDF) or
                  the largest embedded picture (PPTX, which can't be rendered)
  chunk_slides()  [{slide,text,notes}] (+ optional {slide: caption} map) ->
                  [{text, slide, slide_end, idx}], NEVER spans a slide

Why chunks never cross a slide: same reasoning as paper.py's page rule — the
citation locator IS the slide number. Merging slide 4's tail into slide 5's
head would point a citation at the wrong slide for whichever half it didn't
quote.
"""
from __future__ import annotations

import io
from pathlib import Path

from ..config import DECK_CAPTION_MIN_CHARS, DECK_CHUNK_CHARS, DECK_MIN_CHUNK_CHARS, DECK_RENDER_WIDTH
from . import paper as paper_mod
from .paper import _paragraphs, _split_long  # reuse paper's text-hygiene helpers

PPTX_EXT = ".pptx"


# ── Parse ─────────────────────────────────────────────────────────────────────

def parse_deck(path: Path) -> list[dict]:
    """Slide text + speaker notes, 1-indexed. Dispatches on file suffix."""
    if path.suffix.lower() == PPTX_EXT:
        return _parse_pptx(path)
    return _parse_pdf_deck(path)  # default: PDF deck, one page = one slide


def _parse_pdf_deck(path: Path) -> list[dict]:
    pages = paper_mod.parse_pdf(path)  # [{page, text}] — already cleaned
    return [{"slide": p["page"], "text": p["text"], "notes": ""} for p in pages]


def _parse_pptx(path: Path) -> list[dict]:
    from pptx import Presentation

    prs = Presentation(str(path))
    slides: list[dict] = []
    for i, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        _collect_shape_text(slide.shapes, texts)
        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text or ""
        except Exception as exc:  # a malformed notes part shouldn't fail the deck
            print(f"[deck] slide {i}: notes read failed ({type(exc).__name__}: {exc})")
        text = paper_mod._clean("\n\n".join(t for t in texts if t and t.strip()))
        slides.append({"slide": i, "text": text, "notes": paper_mod._clean(notes)})
    return slides


def _collect_shape_text(shapes, out: list[str]) -> None:
    """Recurse into groups and tables; skip pictures (no text to collect)."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                _collect_shape_text(shape.shapes, out)
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text_frame is not None:
                            out.append(cell.text_frame.text)
            elif getattr(shape, "has_text_frame", False):
                out.append(shape.text_frame.text)
        except Exception as exc:  # one odd shape shouldn't sink the whole slide
            print(f"[deck] shape read failed ({type(exc).__name__}: {exc})")


# ── Render (for vision captioning + citation thumbnails) ─────────────────────

def render_slides(path: Path, slides: list[int], *, width: int | None = None) -> dict[int, bytes]:
    """Slide numbers -> representative JPEG. PDF: rasterize the page. PPTX:
    can't be rendered without a graphics engine — use the largest embedded
    picture instead (a slide with no picture and no chart image yields no
    render, which is fine: it just gets no caption)."""
    if not slides:
        return {}
    if path.suffix.lower() == PPTX_EXT:
        return _render_pptx(path, slides, width or DECK_RENDER_WIDTH)
    return _render_pdf(path, slides, width or DECK_RENDER_WIDTH)


def _render_pdf(path: Path, slides: list[int], width: int) -> dict[int, bytes]:
    import pypdfium2 as pdfium

    out: dict[int, bytes] = {}
    doc = pdfium.PdfDocument(str(path))
    try:
        for s in slides:
            idx = s - 1
            if idx < 0 or idx >= len(doc):
                continue
            page = doc[idx]
            try:
                page_width_pt = page.get_size()[0]
                scale = (width / page_width_pt) if page_width_pt else 1.0
                bitmap = page.render(scale=scale)
                try:
                    img = bitmap.to_pil().convert("RGB")
                finally:
                    bitmap.close()
                buf = io.BytesIO()
                img.save(buf, format="JPEG", quality=85)
                out[s] = buf.getvalue()
            except Exception as exc:
                print(f"[deck] slide {s}: render failed ({type(exc).__name__}: {exc})")
            finally:
                page.close()
    finally:
        doc.close()
    return out


def _render_pptx(path: Path, slides: list[int], width: int) -> dict[int, bytes]:
    from PIL import Image
    from pptx import Presentation

    prs = Presentation(str(path))
    slide_list = list(prs.slides)
    out: dict[int, bytes] = {}
    for s in slides:
        idx = s - 1
        if idx < 0 or idx >= len(slide_list):
            continue
        blob = _largest_picture_blob(slide_list[idx].shapes)
        if blob is None:
            continue
        try:
            img = Image.open(io.BytesIO(blob)).convert("RGB")
            if max(img.size) > width:
                img.thumbnail((width, width))
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=85)
            out[s] = buf.getvalue()
        except Exception as exc:
            print(f"[deck] slide {s}: picture decode failed ({type(exc).__name__}: {exc})")
    return out


def resize_jpeg(jpeg: bytes, width: int) -> bytes:
    """Downscale an already-rendered slide JPEG to a smaller width — used to
    turn the (larger) vision-LLM render into a lighter citation thumbnail
    without rasterizing the slide a second time."""
    from PIL import Image

    img = Image.open(io.BytesIO(jpeg))
    if max(img.size) <= width:
        return jpeg
    img = img.convert("RGB")
    img.thumbnail((width, width))
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=85)
    return buf.getvalue()


def _largest_picture_blob(shapes) -> bytes | None:
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    best_area = -1
    best_blob: bytes | None = None
    for shape in shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                blob = _largest_picture_blob(shape.shapes)
                area = shape.width * shape.height if blob else -1
            elif shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
                blob = shape.image.blob
                area = shape.width * shape.height
            else:
                continue
            if blob is not None and area > best_area:
                best_area, best_blob = area, blob
        except Exception:
            continue  # a linked/broken picture with no cached image — skip it
    return best_blob


# ── Chunk (slide-aware — never spans a slide) ────────────────────────────────

def needs_caption(unit: dict) -> bool:
    """Little/no extracted text usually means a chart or diagram drawn as
    vector shapes, not a missing slide — caption it with the vision LLM."""
    return len(unit.get("text", "").strip()) < DECK_CAPTION_MIN_CHARS


def compose_slide_text(unit: dict, caption: str = "") -> str:
    """Slide text + speaker notes + vision caption, in embedding-friendly order."""
    parts = [f"Slide {unit['slide']}"]
    body = (unit.get("text") or "").strip()
    if body:
        parts.append(body)
    notes = (unit.get("notes") or "").strip()
    if notes:
        parts.append(f"[Speaker notes] {notes}")
    if caption:
        parts.append(f"[Slide image] {caption.strip()}")
    return "\n\n".join(parts)


def chunk_slides(units: list[dict], captions: dict[int, str] | None = None, *,
                 chunk_chars: int | None = None, min_chars: int | None = None) -> list[dict]:
    """One chunk per slide normally; a slide whose composed text exceeds
    `chunk_chars` splits at paragraph/sentence boundaries — but a chunk NEVER
    crosses two slides, so its `slide` payload is always exact.

    A whole slide's chunk is NEVER dropped just for being short — a table or
    a one-line agenda slide is still a legitimate citation, and slides are
    short by nature (unlike a paper page, "short" isn't a sign of noise).
    `min_chars` only filters the tiny leftover fragment that splitting one
    OVERSIZED slide into several pieces can produce — e.g. a two-word tail
    after the paragraph before it filled the chunk. A slide with no text, no
    notes, and no caption is skipped entirely (nothing to index).

    Returns [{text, slide, slide_end, idx}]. `slide_end` equals `slide` today
    (chunks don't span slides); the field exists so a future spanning mode
    doesn't need a payload shape change (mirrors paper.py's `page_end`).
    """
    span = chunk_chars or DECK_CHUNK_CHARS
    floor = min_chars if min_chars is not None else DECK_MIN_CHUNK_CHARS
    caption_map = captions or {}

    chunks: list[dict] = []
    for unit in units:
        cap = caption_map.get(unit["slide"], "")
        if not (unit.get("text") or "").strip() and not (unit.get("notes") or "").strip() and not cap:
            continue  # truly blank slide — nothing to index
        text = compose_slide_text(unit, cap)
        paras = _paragraphs(text)
        if not paras:
            continue
        pieces: list[str] = []
        for p in paras:
            pieces.extend(_split_long(p, span) if len(p) > span else [p])
        segments: list[str] = []
        buf = ""
        for piece in pieces:
            candidate = f"{buf}\n\n{piece}" if buf else piece
            if len(candidate) > span and buf:
                segments.append(buf)
                buf = piece
            else:
                buf = candidate
        if buf:
            segments.append(buf)

        multi = len(segments) > 1  # only split slides risk a noise fragment
        for seg in segments:
            seg = seg.strip()
            if multi and len(seg) < floor:
                continue
            chunks.append({"text": seg, "slide": unit["slide"], "slide_end": unit["slide"]})

    for i, c in enumerate(chunks):
        c["idx"] = i
    return chunks


def guess_title(units: list[dict], fallback: str) -> str:
    """First non-trivial line of slide 1 (or the first slide with any text) —
    good enough when the caller didn't supply a title. Slide titles run
    shorter than paper titles, so the floor is lower than paper.guess_title's."""
    for u in units:
        for line in u["text"].splitlines():
            line = line.strip()
            if len(line) >= 3:
                return line[:200]
    return fallback
